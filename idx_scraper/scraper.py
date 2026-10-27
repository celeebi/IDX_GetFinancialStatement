"""
High-level scraper coordinator for IDX financial statements.
Orchestrates API queries, file downloads, parsing, and multi-year compilation.
Supports both online automatic scraping from IDX and completely offline local folder ingestion.
"""

import os
import shutil
import re
import logging
from typing import Dict, Any, List, Optional, Tuple
import pandas as pd
import openpyxl
from .client import IDXClient
from .api import IDXApi
from .excel_parser import parse_financial_statement_file, find_sheet_by_title_or_prefix
from .xbrl_parser import XBRLParser
from .multi_year import merge_statement_tables
from .year_utils import format_year_range_slug, format_year_range_label
from .proforma import (
    FinancialNormalizer, DriverCalculator, BankingDriverCalculator,
    ReadyToForecastExporter, detect_archetype
)

logger = logging.getLogger("idx_scraper.scraper")


def inspect_file_metadata(file_path: str) -> Tuple[Optional[str], Optional[int], Optional[str]]:
    """
    Extract (ticker, year, file_type) from a local file path.
    file_type is 'xlsx' or 'zip' or None.
    """
    fname = os.path.basename(file_path)
    lower_fname = fname.lower()

    file_type = None
    if lower_fname.endswith((".xlsx", ".xls")) and not lower_fname.startswith("~$"):
        file_type = "xlsx"
    elif lower_fname.endswith((".zip", ".xbrl")):
        file_type = "zip"
    else:
        return None, None, None

    ticker = None
    year = None

    # Pattern 1: Standard IDX naming: FinancialStatement-YYYY-Tahunan-TICKER.ext
    m = re.search(r"FinancialStatement-(\d{4})-[A-Za-z0-9_]+-([A-Za-z0-9]+)", fname, re.IGNORECASE)
    if m:
        year = int(m.group(1))
        ticker = m.group(2).upper()
        return ticker, year, file_type

    # Pattern 2: TICKER_YYYY or TICKER-YYYY or TICKER YYYY
    m = re.search(r"\b([A-Za-z]{4})[_\-\s]+(\d{4})\b", fname)
    if m:
        ticker = m.group(1).upper()
        year = int(m.group(2))
        return ticker, year, file_type

    # Pattern 3: YYYY_TICKER or YYYY-TICKER
    m = re.search(r"\b(\d{4})[_\-\s]+([A-Za-z]{4})\b", fname)
    if m:
        year = int(m.group(1))
        ticker = m.group(2).upper()
        return ticker, year, file_type

    # Deep inspection from file content (most reliable for custom filenames)
    if file_type == "xlsx":
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            gi_name = find_sheet_by_title_or_prefix(wb, "general_info")
            if not gi_name and wb.sheetnames:
                gi_name = wb.sheetnames[0]
            if gi_name and gi_name in wb.sheetnames:
                ws = wb[gi_name]
                for row in ws.iter_rows(values_only=True, max_row=45):
                    if not row or len(row) < 2 or row[0] is None:
                        continue
                    k = str(row[0]).strip().lower()
                    v = str(row[1]).strip() if row[1] is not None else ""
                    if not ticker and any(x in k for x in ["kode entitas", "kode emiten", "entity code"]):
                        if v and len(v) <= 6:
                            ticker = v.upper()
                    if not year and any(x in k for x in ["tanggal akhir periode berjalan", "current period end date", "tahun buku"]):
                        my = re.search(r"\b(20\d{2})\b", v)
                        if my:
                            year = int(my.group(1))
            wb.close()
        except Exception as e:
            logger.debug(f"Deep Excel inspection error on {file_path}: {e}")

    elif file_type == "zip":
        try:
            facts, dei, _ = XBRLParser.load_instance(file_path)
            if not ticker:
                ec = dei.get("EntityCode") or dei.get("EntityCentralIndexKey")
                if ec:
                    ticker = ec.upper()
            if not year:
                ed = dei.get("CurrentPeriodEndDate", "")
                my = re.search(r"\b(20\d{2})\b", ed)
                if my:
                    year = int(my.group(1))
        except Exception as e:
            logger.debug(f"Deep XBRL inspection error on {file_path}: {e}")

    # Pattern 4: If still missing, check IMMEDIATE parent or grandparent folder only
    parent = os.path.basename(os.path.dirname(file_path))
    grandparent = os.path.basename(os.path.dirname(os.path.dirname(file_path)))

    if not year and re.match(r"^\d{4}$", parent):
        year = int(parent)
    if not ticker and re.match(r"^[A-Za-z]{4}$", grandparent) and grandparent.upper() not in ("TEMP", "DATA", "TEST", "XLSX", "XBRL", "DOWNLOADS"):
        ticker = grandparent.upper()
    if not ticker and re.match(r"^[A-Za-z]{4}$", parent) and parent.upper() not in ("TEMP", "DATA", "TEST", "XLSX", "XBRL", "DOWNLOADS"):
        ticker = parent.upper()

    return ticker, year, file_type


class IDXFinancialScraper:
    """Coordinator for scraping and compiling IDX financial statements."""

    def __init__(self, output_dir: str = "downloads", client: Optional[IDXClient] = None):
        self.output_dir = output_dir
        self.client = client or IDXClient()
        self.api = IDXApi(self.client)
        os.makedirs(self.output_dir, exist_ok=True)

    def scrape_ticker(
        self,
        ticker: str,
        start_year: Optional[int] = None,
        end_year: Optional[int] = None,
        years: Optional[List[int]] = None,
        period: str = "audit",
        parse_and_export: bool = True
    ) -> Dict[str, Any]:
        """
        Scrape financial statements for a specific ticker across a single year, a range, or an arbitrary list of years.
        
        Args:
            ticker: Stock ticker code (e.g., 'BBCA', 'TLKM', 'ASII')
            start_year: Optional start year
            end_year: Optional end year
            years: Optional explicit list of years (e.g. [2023], [2020, 2021, 2022], [2019, 2021, 2023])
            period: Reporting period ('audit', 'tw1', 'tw2', 'tw3')
            parse_and_export: Whether to parse Excel files and export a multi-year consolidated Excel workbook
        """
        ticker = ticker.strip().upper()

        # Resolve years to query
        if years is not None:
            years_to_query = sorted(list(set(years)))
        elif start_year is not None and end_year is not None:
            years_to_query = list(range(min(start_year, end_year), max(start_year, end_year) + 1))
        elif start_year is not None:
            years_to_query = [start_year]
        else:
            years_to_query = [2020, 2021, 2022, 2023, 2024]

        year_label = format_year_range_label(years_to_query)
        year_slug = format_year_range_slug(years_to_query)
        print(f"\n[+] Searching filings for {ticker} ({year_label}, Period: {period})...")

        # 1. Query filing history
        filings = self.api.get_filing_history(
            kode_emiten=ticker,
            years=years_to_query,
            period=period
        )

        if not filings:
            print(f"[-] No filings found for {ticker} across: {year_label}.")
            return {"ticker": ticker, "status": "no_filings", "downloaded_files": [], "compiled_excel": None}

        # Check if any requested years were not found on IDX
        filing_years_found = {f["year"] for f in filings}
        missing_years = [y for y in years_to_query if y not in filing_years_found]
        if missing_years:
            print(f"  [i] Note: No online filings published on IDX portal for years {missing_years}.")
            print(f"  [+] Found {len(filings)} filings: {sorted(list(filing_years_found))}")
        else:
            print(f"[+] Found {len(filings)} filings for {ticker}.")

        # 2. Download standardized Excel files
        downloaded_files = []
        parsed_reports = []

        ticker_dir = os.path.join(self.output_dir, ticker)
        os.makedirs(ticker_dir, exist_ok=True)

        for f in filings:
            year = f["year"]
            att = f.get("attachment")
            if not att:
                print(f"  [!] Year {year}: No standardized Excel attachment found.")
                continue

            file_path = att.get("File_Path")
            file_name = att.get("File_Name")
            dest_dir = os.path.join(ticker_dir, str(year))
            dest_file = os.path.join(dest_dir, file_name)

            print(f"  [*] Downloading Year {year}: {file_name}...")
            newly_downloaded = self.client.download_file(file_path, dest_file, overwrite=False)
            downloaded_files.append(dest_file)

            # Download XBRL zip if available
            zip_att = f.get("zip_attachment")
            if zip_att and zip_att.get("File_Path"):
                z_path = zip_att.get("File_Path")
                z_name = zip_att.get("File_Name") or "instance.zip"
                dest_zip = os.path.join(dest_dir, z_name)
                print(f"  [*] Downloading Year {year} XBRL: {z_name}...")
                self.client.download_file(z_path, dest_zip, overwrite=False)
                downloaded_files.append(dest_zip)

            # 3. Parse if requested
            if parse_and_export and os.path.exists(dest_file):
                try:
                    parsed = parse_financial_statement_file(dest_file)
                    parsed["year"] = year
                    parsed_reports.append(parsed)
                except Exception as e:
                    logger.error(f"Error parsing {dest_file}: {e}")

        compiled_excel_path = None
        if parse_and_export and parsed_reports:
            print(f"[+] Compiling multi-year statements ({len(parsed_reports)} years)...")
            
            # Sort reports chronologically
            parsed_reports.sort(key=lambda x: x.get("year", 0))

            bs_multi = merge_statement_tables(parsed_reports, "balance_sheet")
            is_multi = merge_statement_tables(parsed_reports, "income_statement")
            cf_multi = merge_statement_tables(parsed_reports, "cash_flow")

            # Collect latest general info and add author credit
            latest_info = parsed_reports[-1].get("general_info", {})
            latest_info["Program Author / Developer"] = "abas YP74A"
            info_df = pd.DataFrame(list(latest_info.items()), columns=["Field", "Value"])

            compiled_excel_name = f"{ticker}_financial_statements_{year_slug}.xlsx"
            compiled_excel_path = os.path.join(ticker_dir, compiled_excel_name)

            with pd.ExcelWriter(compiled_excel_path, engine="openpyxl") as writer:
                if not bs_multi.empty:
                    bs_multi.to_excel(writer, sheet_name="Balance_Sheet", index=False)
                if not is_multi.empty:
                    is_multi.to_excel(writer, sheet_name="Income_Statement", index=False)
                if not cf_multi.empty:
                    cf_multi.to_excel(writer, sheet_name="Cash_Flow", index=False)
                if not info_df.empty:
                    info_df.to_excel(writer, sheet_name="General_Info", index=False)
                writer.book.properties.creator = "abas YP74A"
                writer.book.properties.lastModifiedBy = "abas YP74A"
                writer.book.properties.description = "Compiled IDX Financial Statements by abas YP74A"

            print(f"[OK] Successfully compiled multi-year workbook: {compiled_excel_path}")

        return {
            "ticker": ticker,
            "status": "success",
            "filings_count": len(filings),
            "downloaded_files": downloaded_files,
            "compiled_excel": compiled_excel_path,
            "parsed_reports": parsed_reports
        }

    def _compile_and_export_ready_model(
        self,
        ticker: str,
        valid_years: List[int],
        compiled_excel: Optional[str] = None,
        info_dict: Optional[Dict[str, str]] = None,
        copy_to_dir: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Internal coordinator to normalize multi-year statements, compute operational drivers,
        and export the standardized Ready-to-Forecast Excel model.
        Used by both online scraping and offline local folder ingestion.
        """
        ticker = ticker.strip().upper()
        ticker_dir = os.path.join(self.output_dir, ticker)
        year_slug = format_year_range_slug(valid_years)

        df_bs = pd.DataFrame()
        df_is = pd.DataFrame()
        df_cf = pd.DataFrame()
        sheet_names = []

        if info_dict is None:
            info_dict = {}

        if compiled_excel and os.path.exists(compiled_excel):
            try:
                xls = pd.ExcelFile(compiled_excel)
                sheet_names = xls.sheet_names
                if "Balance_Sheet" in sheet_names:
                    df_bs = pd.read_excel(xls, "Balance_Sheet")
                if "Income_Statement" in sheet_names:
                    df_is = pd.read_excel(xls, "Income_Statement")
                if "Cash_Flow" in sheet_names:
                    df_cf = pd.read_excel(xls, "Cash_Flow")
                if "General_Info" in sheet_names and not info_dict:
                    df_info = pd.read_excel(xls, "General_Info")
                    for _, r in df_info.iterrows():
                        if pd.notna(r.iloc[0]):
                            info_dict[str(r.iloc[0])] = str(r.iloc[1]) if pd.notna(r.iloc[1]) else ""
            except Exception as e:
                logger.warning(f"Could not load compiled tables from {compiled_excel}: {e}")

        # Detect sector & financial modeling archetype
        archetype = detect_archetype(
            ticker=ticker,
            general_info=info_dict,
            sheet_names=sheet_names
        )
        arch_desc = "Banking & Financial Institutions Model" if archetype == "banking" else "Corporate Operating Model"
        print(f"[+] Detected Archetype: {archetype.upper()} ({arch_desc})")

        normalizer = FinancialNormalizer(target_scale="millions")
        is_dict = {}
        bs_dict = {}
        cf_dict = {}

        print(f"[+] Normalizing statements across {len(valid_years)} year(s) (Primary: XBRL, Fallback: Excel)...")

        for yr in valid_years:
            year_dir = os.path.join(ticker_dir, str(yr))
            zip_candidates = [
                os.path.join(year_dir, "instance.zip"),
                os.path.join(year_dir, f"instance-{yr}.zip"),
            ]
            if os.path.exists(year_dir):
                for f_name in os.listdir(year_dir):
                    if f_name.endswith(".zip"):
                        candidate = os.path.join(year_dir, f_name)
                        if candidate not in zip_candidates:
                            zip_candidates.append(candidate)

            xbrl_loaded = False
            for z_path in zip_candidates:
                if os.path.exists(z_path):
                    try:
                        facts, dei, contexts = XBRLParser.load_instance(z_path)
                        if facts:
                            if dei and not info_dict:
                                info_dict.update(dei)
                            # Re-verify archetype with XBRL facts
                            detected_arch = detect_archetype(ticker=ticker, general_info=info_dict, facts=facts)
                            if detected_arch != archetype:
                                archetype = detected_arch

                            if archetype == "banking":
                                is_stmt, bs_stmt, cf_stmt = XBRLParser.extract_banking_statements(facts, yr, dei=dei)
                            else:
                                is_stmt, bs_stmt, cf_stmt = XBRLParser.extract_normalized_statements(facts, yr, dei=dei)
                            is_dict[yr] = is_stmt
                            bs_dict[yr] = bs_stmt
                            cf_dict[yr] = cf_stmt
                            xbrl_loaded = True
                            print(f"  [+] Year {yr}: Loaded directly from primary XBRL data (BS Balance Check: {bs_stmt.balance_check:0.2f})")
                            break
                    except Exception as e:
                        logger.warning(f"Failed to load XBRL for {ticker} {yr} from {z_path}: {e}")

            if not xbrl_loaded:
                # Check if next year's ZIP exists and can provide prior-year comparative data
                next_yr = yr + 1
                next_yr_dir = os.path.join(ticker_dir, str(next_yr))
                next_zip_candidates = [
                    os.path.join(next_yr_dir, "instance.zip"),
                    os.path.join(next_yr_dir, f"instance-{next_yr}.zip"),
                ]
                for next_z in next_zip_candidates:
                    if os.path.exists(next_z):
                        try:
                            facts_next, dei_next, _ = XBRLParser.load_instance(next_z)
                            if facts_next:
                                if dei_next and not info_dict:
                                    info_dict.update(dei_next)
                                if archetype == "banking":
                                    is_stmt, bs_stmt, cf_stmt = XBRLParser.extract_banking_statements(facts_next, yr, dei=dei_next, period_target="prior")
                                else:
                                    is_stmt, bs_stmt, cf_stmt = XBRLParser.extract_normalized_statements(facts_next, yr, dei=dei_next, period_target="prior")
                                if bs_stmt.total_assets > 0:
                                    is_dict[yr] = is_stmt
                                    bs_dict[yr] = bs_stmt
                                    cf_dict[yr] = cf_stmt
                                    xbrl_loaded = True
                                    print(f"  [+] Year {yr}: Extracted from {next_yr} comparative XBRL data (BS Balance Check: {bs_stmt.balance_check:0.2f})")
                                    break
                        except Exception as e:
                            logger.debug(f"Failed prior-period comparative extraction from {next_z}: {e}")

            if not xbrl_loaded:
                print(f"  [*] Year {yr}: Falling back to Excel waterfall resolution...")
                if archetype == "banking":
                    is_stmt, bs_stmt, cf_stmt = normalizer.normalize_banking_year(
                        year=yr,
                        df_is=df_is,
                        df_bs=df_bs,
                        df_cf=df_cf,
                        general_info=info_dict
                    )
                else:
                    is_stmt, bs_stmt, cf_stmt = normalizer.normalize_year(
                        year=yr,
                        df_is=df_is,
                        df_bs=df_bs,
                        df_cf=df_cf,
                        general_info=info_dict
                    )
                is_dict[yr] = is_stmt
                bs_dict[yr] = bs_stmt
                cf_dict[yr] = cf_stmt

        # 4. Compute operational drivers
        if archetype == "banking":
            print("[+] Calculating historical banking drivers (NIM, CoF, LDR, CASA, BOPO/CIR, Credit Cost, ROA, ROE)...")
            drivers_by_year = BankingDriverCalculator.calculate_drivers(valid_years, is_dict, bs_dict, cf_dict)
            benchmarks = BankingDriverCalculator.build_summary_benchmarks(drivers_by_year)
        else:
            print("[+] Calculating historical forecasting drivers (DSO, DIO, DPO, margins, Capex)...")
            drivers_by_year = DriverCalculator.calculate_drivers(valid_years, is_dict, bs_dict, cf_dict)
            benchmarks = DriverCalculator.build_summary_benchmarks(drivers_by_year)

        # 5. Export 5-tab ready-to-forecast workbook
        output_name = f"{ticker}_ReadyToForecast_Financials_{year_slug}.xlsx"
        output_path = os.path.join(ticker_dir, output_name)

        print(f"[+] Exporting Ready-to-Forecast Excel model: {output_name}...")
        ReadyToForecastExporter.export(
            dest_path=output_path,
            ticker=ticker,
            years=valid_years,
            is_dict=is_dict,
            bs_dict=bs_dict,
            cf_dict=cf_dict,
            drivers_by_year=drivers_by_year,
            general_info=info_dict,
            archetype=archetype
        )

        print(f"[OK] Ready-to-Forecast model exported successfully to:\n     {os.path.abspath(output_path)}")

        # Copy to input folder if requested and distinct from output dir
        if copy_to_dir and os.path.isdir(copy_to_dir):
            if os.path.abspath(copy_to_dir) != os.path.abspath(ticker_dir):
                target_copy = os.path.join(copy_to_dir, output_name)
                try:
                    shutil.copy2(output_path, target_copy)
                    print(f"[+] Also saved a copy to input folder:\n     {os.path.abspath(target_copy)}")
                except Exception as e:
                    logger.debug(f"Could not copy to input folder: {e}")

        return {
            "ticker": ticker,
            "status": "success",
            "ready_excel": output_path,
            "archetype": archetype,
            "years": valid_years,
            "is_dict": is_dict,
            "bs_dict": bs_dict,
            "cf_dict": cf_dict,
            "drivers_by_year": drivers_by_year,
            "benchmarks": benchmarks
        }

    def build_ready_to_forecast(
        self,
        ticker: str,
        start_year: Optional[int] = None,
        end_year: Optional[int] = None,
        years: Optional[List[int]] = None
    ) -> Dict[str, Any]:
        """
        Extract, normalize, and compile clean historical financial statements
        and operational drivers into a Ready-to-Forecast Excel model across any time horizon.
        """
        ticker = ticker.strip().upper()

        # Resolve years to process
        if years is not None:
            years_to_process = sorted(list(set(years)))
        elif start_year is not None and end_year is not None:
            years_to_process = list(range(min(start_year, end_year), max(start_year, end_year) + 1))
        elif start_year is not None:
            years_to_process = [start_year]
        else:
            years_to_process = [2020, 2021, 2022, 2023, 2024]

        year_label = format_year_range_label(years_to_process)
        print(f"\n[+] Preparing Ready-to-Forecast Financials for {ticker} ({year_label})...")

        # 1. Ensure files are scraped and multi-year raw tables exist
        scrape_res = self.scrape_ticker(
            ticker=ticker,
            years=years_to_process,
            period="audit",
            parse_and_export=True
        )

        compiled_excel = scrape_res.get("compiled_excel")
        if not compiled_excel or not os.path.exists(compiled_excel):
            print(f"[-] Could not find compiled historical statements for {ticker}.")
            return {"status": "error", "message": "Raw files could not be scraped"}

        ticker_dir = os.path.join(self.output_dir, ticker)

        # Filter years_to_process to only those years where data was actually found
        xls = pd.ExcelFile(compiled_excel)
        df_bs = pd.read_excel(xls, "Balance_Sheet") if "Balance_Sheet" in xls.sheet_names else pd.DataFrame()
        df_is = pd.read_excel(xls, "Income_Statement") if "Income_Statement" in xls.sheet_names else pd.DataFrame()

        valid_years = []
        for y in years_to_process:
            y_dir = os.path.join(ticker_dir, str(y))
            if os.path.exists(y_dir) or str(y) in df_bs.columns or str(y) in df_is.columns:
                valid_years.append(y)

        if not valid_years:
            valid_years = years_to_process

        return self._compile_and_export_ready_model(
            ticker=ticker,
            valid_years=valid_years,
            compiled_excel=compiled_excel
        )

    def process_local_folder(
        self,
        folder_path: str,
        output_dir: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        """
        Process an arbitrary local folder containing downloaded financial statement
        files (.xlsx and/or .zip XBRL packages) without any network calls to IDX.
        
        Supports flat directories, year subfolders, and multi-company folders.
        Automatically discovers tickers, years, and archetypes from files.

        Args:
            folder_path: Path to directory or file containing downloaded financial statements.
            output_dir: Optional custom output directory (defaults to self.output_dir).

        Returns:
            List of result dictionaries for each company processed.
        """
        clean_path = os.path.abspath(os.path.expanduser(folder_path.strip().strip('"').strip("'")))
        if not os.path.exists(clean_path):
            print(f"[-] Directory or file does not exist: {clean_path}")
            return []

        if output_dir:
            self.output_dir = output_dir
            os.makedirs(self.output_dir, exist_ok=True)

        # 1. Discover all candidate files
        discovered_files: List[str] = []
        if os.path.isfile(clean_path):
            discovered_files.append(clean_path)
            scan_root = os.path.dirname(clean_path)
        else:
            scan_root = clean_path
            for root, _, files in os.walk(clean_path):
                for f in files:
                    if f.startswith("~$") or "ReadyToForecast" in f or "_financial_statements_" in f:
                        continue
                    ext = os.path.splitext(f)[1].lower()
                    if ext in (".xlsx", ".xls", ".zip", ".xbrl"):
                        discovered_files.append(os.path.join(root, f))

        if not discovered_files:
            print(f"[-] No financial statement files (.xlsx or .zip) found in: {clean_path}")
            return []

        print(f"\n[+] Scanning {len(discovered_files)} file(s) in: {clean_path}...")

        # 2. Extract metadata and group by ticker and year
        # companies[ticker][year] = {"xlsx": [...], "zip": [...]}
        companies: Dict[str, Dict[int, Dict[str, List[str]]]] = {}

        for fpath in discovered_files:
            ticker, year, ftype = inspect_file_metadata(fpath)
            if ticker and year and ftype:
                if ticker not in companies:
                    companies[ticker] = {}
                if year not in companies[ticker]:
                    companies[ticker][year] = {"xlsx": [], "zip": []}
                companies[ticker][year][ftype].append(fpath)
                logger.debug(f"Detected: {ticker} {year} [{ftype}] -> {os.path.basename(fpath)}")
            else:
                logger.debug(f"Skipping unclassified file: {fpath}")

        if not companies:
            print(f"[-] Could not identify valid stock ticker or reporting years from files in: {clean_path}")
            return []

        print(f"[+] Discovered {len(companies)} company/companies in local files:")
        for t, y_dict in sorted(companies.items()):
            print(f"    • {t}: {sorted(y_dict.keys())}")

        # 3. Process each discovered company
        results: List[Dict[str, Any]] = []

        for ticker, yr_dict in sorted(companies.items()):
            valid_years = sorted(yr_dict.keys())
            year_slug = format_year_range_slug(valid_years)
            year_label = format_year_range_label(valid_years)

            print(f"\n" + "=" * 74)
            print(f"      PROCESSING LOCAL FILES FOR {ticker} ({year_label})")
            print("=" * 74)

            ticker_dir = os.path.join(self.output_dir, ticker)
            os.makedirs(ticker_dir, exist_ok=True)

            parsed_reports = []
            info_dict = {}

            # Stage files into output_dir/<ticker>/<year>/
            for yr in valid_years:
                dest_yr_dir = os.path.join(ticker_dir, str(yr))
                os.makedirs(dest_yr_dir, exist_ok=True)

                # Copy/stage ZIP files
                for zf in yr_dict[yr]["zip"]:
                    dest_z = os.path.join(dest_yr_dir, "instance.zip")
                    if os.path.abspath(zf) != os.path.abspath(dest_z):
                        try:
                            shutil.copy2(zf, dest_z)
                        except Exception as e:
                            logger.debug(f"Copying {zf} -> {dest_z}: {e}")

                # Copy/stage XLSX files & parse
                for xf in yr_dict[yr]["xlsx"]:
                    dest_x = os.path.join(dest_yr_dir, os.path.basename(xf))
                    if os.path.abspath(xf) != os.path.abspath(dest_x):
                        try:
                            shutil.copy2(xf, dest_x)
                        except Exception as e:
                            logger.debug(f"Copying {xf} -> {dest_x}: {e}")
                    try:
                        parsed = parse_financial_statement_file(xf)
                        parsed["year"] = yr
                        parsed_reports.append(parsed)
                        if "general_info" in parsed and parsed["general_info"]:
                            info_dict.update(parsed["general_info"])
                    except Exception as e:
                        logger.warning(f"Error parsing local Excel file {xf}: {e}")

            # Compile multi-year raw tables if Excel reports were parsed
            compiled_excel_path = None
            if parsed_reports:
                print(f"[+] Compiling multi-year statements ({len(parsed_reports)} report(s))...")
                parsed_reports.sort(key=lambda x: x.get("year", 0))

                bs_multi = merge_statement_tables(parsed_reports, "balance_sheet")
                is_multi = merge_statement_tables(parsed_reports, "income_statement")
                cf_multi = merge_statement_tables(parsed_reports, "cash_flow")

                latest_info = parsed_reports[-1].get("general_info", {})
                latest_info["Program Author / Developer"] = "abas YP74A"
                info_df = pd.DataFrame(list(latest_info.items()), columns=["Field", "Value"])

                compiled_excel_name = f"{ticker}_financial_statements_{year_slug}.xlsx"
                compiled_excel_path = os.path.join(ticker_dir, compiled_excel_name)

                with pd.ExcelWriter(compiled_excel_path, engine="openpyxl") as writer:
                    if not bs_multi.empty:
                        bs_multi.to_excel(writer, sheet_name="Balance_Sheet", index=False)
                    if not is_multi.empty:
                        is_multi.to_excel(writer, sheet_name="Income_Statement", index=False)
                    if not cf_multi.empty:
                        cf_multi.to_excel(writer, sheet_name="Cash_Flow", index=False)
                    if not info_df.empty:
                        info_df.to_excel(writer, sheet_name="General_Info", index=False)
                    writer.book.properties.creator = "abas YP74A"
                    writer.book.properties.lastModifiedBy = "abas YP74A"
                    writer.book.properties.description = "Compiled IDX Financial Statements by abas YP74A"

                print(f"[OK] Successfully compiled raw multi-year workbook: {compiled_excel_path}")
                if os.path.abspath(scan_root) != os.path.abspath(ticker_dir):
                    try:
                        shutil.copy2(compiled_excel_path, os.path.join(scan_root, compiled_excel_name))
                    except Exception:
                        pass

            # Build and export Ready-to-Forecast Model
            res = self._compile_and_export_ready_model(
                ticker=ticker,
                valid_years=valid_years,
                compiled_excel=compiled_excel_path,
                info_dict=info_dict if info_dict else None,
                copy_to_dir=scan_root if os.path.abspath(scan_root) != os.path.abspath(ticker_dir) else None
            )
            results.append(res)

        return results
