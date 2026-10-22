"""
Parser for standardized IDX financial statement Excel files.
Extracts General Information, Balance Sheet, Income Statement, and Cash Flows.
"""

import os
import re
import logging
from typing import Dict, Any, Optional, List, Tuple
import openpyxl
import pandas as pd

logger = logging.getLogger("idx_scraper.excel_parser")

# Sheet prefix mappings according to all IDX XBRL taxonomies:
# - General Industry: 1210000 (BS), 1310000/1321000 (IS), 1510000 (CF)
# - Infrastructure/Trade/Services: 3210000 (BS), 3312000 (IS), 3510000 (CF)
# - Financial & Sharia: 4220000 (BS), 4322000 (IS), 4510000/4610000 (CF)
SHEET_MAPPINGS = {
    "general_info": ["1000000", "1110000"],
    "balance_sheet": ["1210000", "2100000", "3210000", "4220000", "121", "210", "321", "422"],
    "income_statement": ["1310000", "1321000", "3100000", "3312000", "4322000", "131", "132", "310", "331", "432"],
    "cash_flow": ["1510000", "1520000", "3510000", "3520000", "4510000", "4610000", "151", "152", "351", "352", "451", "461"],
}


def find_sheet_by_title_or_prefix(wb, target_type: str) -> Optional[str]:
    """
    Find sheet by inspecting row 0 title first (most robust), falling back to taxonomy code prefixes.
    """
    title_keywords = {
        "balance_sheet": ["statement of financial position", "posisi keuangan", "neraca"],
        "income_statement": ["statement of profit or loss", "laba rugi"],
        "cash_flow": ["statement of cash flows", "arus kas"],
        "general_info": ["general information", "informasi umum"]
    }

    keywords = title_keywords.get(target_type, [])

    # Pass 1: Check row 0 title in each sheet
    for name in wb.sheetnames:
        # Quick skip hidden or metadata
        if name.lower() in ("token", "hidden", "inlinexbrl", "context"):
            continue
        try:
            ws = wb[name]
            r0 = list(ws.iter_rows(values_only=True, max_row=1))
            t = str(r0[0][0]).lower() if r0 and r0[0] and r0[0][0] is not None else ""
            if any(k in t for k in keywords):
                return name
        except Exception:
            pass

    # Pass 2: Prefix fallback
    prefixes = SHEET_MAPPINGS.get(target_type, [])
    for p in prefixes:
        for name in wb.sheetnames:
            if name.strip().startswith(p):
                return name

    return None


def parse_general_info(ws) -> Dict[str, str]:
    """Parse sheet 1000000 to extract entity profile."""
    info = {}
    rows = list(ws.iter_rows(values_only=True))
    for row in rows:
        if not row or len(row) < 2 or row[0] is None:
            continue
        key = str(row[0]).strip()
        val = str(row[1]).strip() if row[1] is not None else ""
        if key and not key.startswith("[1000000]"):
            info[key] = val
    return info


def parse_statement_sheet(ws, report_year: Optional[int] = None) -> Tuple[str, Dict[str, str], pd.DataFrame]:
    """
    Parse a tabular financial statement sheet into a structured DataFrame.
    Returns: (sheet_title, dates_dict, DataFrame)
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return "", {}, pd.DataFrame()

    title = str(rows[0][0]) if rows and rows[0] and rows[0][0] is not None else ""
    date_col1: Optional[Tuple[int, str]] = None
    date_col2: Optional[Tuple[int, str]] = None
    data_start_idx = 0

    # Search top 10 rows for date columns or XBRL context headers (e.g., CurrentYearInstant)
    for idx, row in enumerate(rows[:10]):
        row_str = " ".join(str(c) for c in row if c is not None)
        
        # Check for XBRL context names
        if "currentyear" in row_str.lower():
            data_start_idx = idx + 1
            for c_idx, c in enumerate(row):
                if not c:
                    continue
                c_lower = str(c).lower()
                if "current" in c_lower and date_col1 is None:
                    h_val = f"{report_year}" if report_year else str(c).strip()
                    date_col1 = (c_idx, h_val)
                elif "prior" in c_lower and date_col2 is None:
                    h_val = f"{report_year - 1}" if report_year else str(c).strip()
                    date_col2 = (c_idx, h_val)
            break

        # Check for date strings like '31 December 2023' or '31 Mar'
        elif re.search(r'\d{4}', row_str) and any(m in row_str.lower() for m in ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "des"]):
            data_start_idx = idx + 1
            date_indices = [c_idx for c_idx, c in enumerate(row) if c and re.search(r'\d{4}', str(c))]
            if len(date_indices) >= 1:
                date_col1 = (date_indices[0], str(row[date_indices[0]]).strip())
            if len(date_indices) >= 2:
                date_col2 = (date_indices[1], str(row[date_indices[1]]).strip())
            break

    records = []
    for row in rows[data_start_idx:]:
        if not row or not any(row):
            continue
        label_id = str(row[0]).strip() if row[0] is not None else ""
        if not label_id:
            continue

        val1 = None
        val2 = None
        label_en = ""

        if date_col1 and date_col1[0] < len(row):
            val1 = row[date_col1[0]]
        if date_col2 and date_col2[0] < len(row):
            val2 = row[date_col2[0]]

        # English label is in column 3 (or 2 if only 1 period)
        if len(row) > 3 and row[3] is not None and isinstance(row[3], str):
            label_en = str(row[3]).strip()
        elif len(row) > 2 and row[2] is not None and isinstance(row[2], str) and not date_col2:
            label_en = str(row[2]).strip()

        # Clean numerical values
        def clean_val(v):
            if v is None:
                return None
            if isinstance(v, (int, float)):
                return float(v)
            try:
                # Remove commas and spaces
                v_str = str(v).replace(",", "").strip()
                if v_str in ("", "-", "n/a", "N/A"):
                    return None
                return float(v_str)
            except (ValueError, TypeError):
                return None

        records.append({
            "Item_ID": label_id,
            "Item_EN": label_en,
            "Current_Value": clean_val(val1),
            "Prior_Value": clean_val(val2),
        })

    cur_header = date_col1[1] if date_col1 else "Current_Period"
    prior_header = date_col2[1] if date_col2 else "Prior_Period"

    df = pd.DataFrame(records)
    if not df.empty:
        # Rename value columns to exact dates if available
        df.rename(columns={"Current_Value": cur_header, "Prior_Value": prior_header}, inplace=True)

    dates = {"current": cur_header, "prior": prior_header}
    return title, dates, df


def parse_financial_statement_file(file_path: str, report_year: Optional[int] = None) -> Dict[str, Any]:
    """
    Parse a complete IDX standardized Excel file.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    # Auto-detect year from filename or path if not provided
    if report_year is None:
        m = re.search(r'\b(20\d{2})\b', os.path.basename(file_path)) or re.search(r'[\\/](20\d{2})[\\/]', file_path)
        if m:
            report_year = int(m.group(1))

    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    sheetnames = wb.sheetnames

    res: Dict[str, Any] = {
        "file_path": file_path,
        "year": report_year,
        "general_info": {},
        "balance_sheet": pd.DataFrame(),
        "income_statement": pd.DataFrame(),
        "cash_flow": pd.DataFrame(),
        "dates": {},
        "sheets_found": {}
    }

    # 1. General Info
    info_sheet = find_sheet_by_title_or_prefix(wb, "general_info")
    if info_sheet:
        res["sheets_found"]["general_info"] = info_sheet
        res["general_info"] = parse_general_info(wb[info_sheet])

    # 2. Balance Sheet
    bs_sheet = find_sheet_by_title_or_prefix(wb, "balance_sheet")
    if bs_sheet:
        res["sheets_found"]["balance_sheet"] = bs_sheet
        title, dates, df = parse_statement_sheet(wb[bs_sheet], report_year=report_year)
        res["balance_sheet"] = df
        res["dates"]["balance_sheet"] = dates

    # 3. Income Statement
    is_sheet = find_sheet_by_title_or_prefix(wb, "income_statement")
    if is_sheet:
        res["sheets_found"]["income_statement"] = is_sheet
        title, dates, df = parse_statement_sheet(wb[is_sheet], report_year=report_year)
        res["income_statement"] = df
        res["dates"]["income_statement"] = dates

    # 4. Cash Flow
    cf_sheet = find_sheet_by_title_or_prefix(wb, "cash_flow")
    if cf_sheet:
        res["sheets_found"]["cash_flow"] = cf_sheet
        title, dates, df = parse_statement_sheet(wb[cf_sheet], report_year=report_year)
        res["cash_flow"] = df
        res["dates"]["cash_flow"] = dates

    wb.close()
    return res
