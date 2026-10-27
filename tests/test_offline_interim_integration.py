"""
Integration test for offline folder ingestion with mixed annual and interim quarter files.
"""

import os
import sys
import shutil
import tempfile
import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from idx_scraper.scraper import IDXFinancialScraper

def run_test():
    tmp_dir = tempfile.mkdtemp(prefix="idx_test_interim_")
    out_dir = os.path.join(tmp_dir, "output")
    os.makedirs(out_dir, exist_ok=True)
    
    try:
        # Source files
        bris_2022_dir = os.path.join("downloads", "BRIS", "2022")
        bris_2023_dir = os.path.join("downloads", "BRIS", "2023")

        # 1. Stage annual 2022 & 2023
        f_2022_xlsx = os.path.join(bris_2022_dir, "FinancialStatement-2022-Tahunan-BRIS.xlsx")
        f_2022_zip = os.path.join(bris_2022_dir, "instance.zip")
        f_2023_xlsx = os.path.join(bris_2023_dir, "FinancialStatement-2023-Tahunan-BRIS.xlsx")
        f_2023_zip = os.path.join(bris_2023_dir, "instance.zip")

        shutil.copy2(f_2022_xlsx, os.path.join(tmp_dir, "FinancialStatement-2022-Tahunan-BRIS.xlsx"))
        shutil.copy2(f_2022_zip, os.path.join(tmp_dir, "instance-2022.zip"))
        shutil.copy2(f_2023_xlsx, os.path.join(tmp_dir, "FinancialStatement-2023-Tahunan-BRIS.xlsx"))
        shutil.copy2(f_2023_zip, os.path.join(tmp_dir, "instance-2023.zip"))

        # 2. Stage interim 2024-tw2 by copying 2023 files named as 2024-tw2
        shutil.copy2(f_2023_xlsx, os.path.join(tmp_dir, "FinancialStatement-2024-tw2-BRIS.xlsx"))
        shutil.copy2(f_2023_zip, os.path.join(tmp_dir, "FinancialStatement-2024-tw2-BRIS.zip"))

        print(f"[+] Running local folder ingestion test on: {tmp_dir}")
        scraper = IDXFinancialScraper(output_dir=out_dir)
        results = scraper.process_local_folder(tmp_dir)

        assert len(results) == 1, f"Expected 1 company result, got {len(results)}"
        res = results[0]
        assert res["ticker"] == "BRIS", f"Expected BRIS, got {res['ticker']}"
        assert res["status"] == "success", f"Expected success, got {res['status']}"
        
        # Verify discovered periods
        expected_periods = ["2022", "2023", "2024-Q2"]
        assert res["years"] == expected_periods, f"Expected {expected_periods}, got {res['years']}"
        print(f"[OK] Periods discovered correctly: {res['years']}")

        # Verify balance check
        for p in expected_periods:
            bs = res["bs_dict"][p]
            assert abs(bs.balance_check) < 0.05, f"Period {p} balance check failed: {bs.balance_check}"
            print(f"  [OK] Balance check for {p}: {bs.balance_check:.2f}")

        # Verify exported workbook
        ready_excel = res["ready_excel"]
        assert os.path.exists(ready_excel), f"Workbook does not exist: {ready_excel}"
        wb = openpyxl.load_workbook(ready_excel)
        
        # Check author credit
        assert wb.properties.creator == "abas YP74A", f"Creator mismatch: {wb.properties.creator}"
        print(f"[OK] Workbook author metadata: {wb.properties.creator}")

        # Check column headers on Balance Sheet
        ws_bs = wb["Historical_BS"]
        headers_bs = [ws_bs.cell(row=4, column=col).value for col in range(2, 5)]
        expected_headers = ["2022", "2023", "2024-Q2 (6M)"]
        assert headers_bs == expected_headers, f"Expected headers {expected_headers}, got {headers_bs}"
        print(f"[OK] Balance sheet column headers: {headers_bs}")

        # Check Operational Drivers tab headers
        ws_drv = wb["Forecasting_Drivers"]
        drv_headers = [ws_drv.cell(row=4, column=col).value for col in range(2, 5)]
        assert drv_headers == expected_headers, f"Expected drv headers {expected_headers}, got {drv_headers}"
        print(f"[OK] Drivers column headers: {drv_headers}")

        wb.close()
        print("\n[ALL TESTS PASSED SUCCESSFULLY!]")

    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

if __name__ == "__main__":
    run_test()
